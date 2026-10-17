"""Private child entry point: selected Office parser only, never generic dispatch."""
import os
import sys
import unicodedata


def _has_content(value):
    return value is not None and any(
        not char.isspace() and unicodedata.category(char) not in ('Cc', 'Cf')
        for char in str(value)
    )


def _deny_external(event, arguments):
    if event in {'socket.connect', 'socket.connect_ex', 'socket.getaddrinfo',
                 'socket.bind', 'subprocess.Popen', 'os.system',
                 'os.posix_spawn', 'os.posix_spawnp', 'os.fork'}:
        raise RuntimeError('external_operation_denied')


def _xlsx_has_content(source):
    import openpyxl
    workbook = openpyxl.load_workbook(
        source, read_only=True, data_only=True, keep_links=False,
    )
    try:
        return any(_has_content(value) for sheet in workbook
                   for row in sheet.iter_rows(values_only=True) for value in row)
    finally:
        workbook.close()


def _pptx_has_content(source):
    import pptx
    presentation = pptx.Presentation(source)

    def shape_content(shape):
        if shape.has_text_frame and _has_content(shape.text):
            return True
        if shape.has_table and any(_has_content(cell.text)
                                   for row in shape.table.rows for cell in row.cells):
            return True
        if shape.has_chart and any(_has_content(value)
                                  for series in shape.chart.series for value in series.values):
            return True
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            return any(shape_content(child) for child in shape.shapes)
        return False

    content = False
    for slide in presentation.slides:
        slide_content = any(shape_content(shape) for shape in slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            slide_content = slide_content or (frame is not None and _has_content(frame.text))
        if not slide_content and len(slide.shapes):
            # A picture/diagram or empty placeholders cannot prove readable text.
            raise ValueError('unreadable_slide')
        content = content or slide_content
    return content


def convert(extension, source):
    from markitdown import StreamInfo
    if extension == '.xlsx':
        from markitdown.converters import XlsxConverter
        converter = XlsxConverter()
        has_content = _xlsx_has_content(source)
    elif extension == '.pptx':
        from markitdown.converters import PptxConverter
        converter = PptxConverter()
        has_content = _pptx_has_content(source)
    else:
        raise ValueError('unsupported_format')
    if not has_content:
        raise ValueError('empty_document')
    with open(source, 'rb') as stream:
        result = converter.convert(stream, StreamInfo(extension=extension))
    if not _has_content(result.markdown):
        raise ValueError('empty_document')
    return result.markdown


def main():
    if len(sys.argv) != 3 or sys.argv[1] not in ('.xlsx', '.pptx'):
        return 2
    if sys.version_info < (3, 10) or not os.path.isabs(sys.argv[2]):
        return 2
    sys.addaudithook(_deny_external)
    try:
        text = convert(sys.argv[1], sys.argv[2])
        sys.stdout.buffer.write(text.encode('utf-8'))
        return 0
    except Exception:
        sys.stderr.write('office_parser_failed\n')
        return 2


if __name__ == '__main__':
    sys.exit(main())
